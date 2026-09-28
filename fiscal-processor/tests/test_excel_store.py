from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from fiscal_processor.adapters.excel.openpyxl_store import (
    OpenpyxlInvoiceStore,
    UpsertOutcome,
    WorkbookContractError,
    WorkbookLockedError,
)
from fiscal_processor.domain.invoice import FiscalExtraction
from fiscal_processor.domain.status import (
    DocumentType,
    ExtractionMode,
    ProcessingStatus,
    QualityFlag,
)

NOW = datetime(2026, 9, 27, 21, 0, tzinfo=UTC)


def extraction(
    *,
    sha: str = "a" * 64,
    filename: str = "nf.pdf",
    invoice: str = "004.241.885",
    amount: str = "112000.00",
    status: ProcessingStatus = ProcessingStatus.OK,
) -> FiscalExtraction:
    return FiscalExtraction(
        source_sha256=sha,
        source_filename=filename,
        document_type=DocumentType.NFE,
        extraction_mode=ExtractionMode.NATIVE_TEXT,
        status=status,
        invoice_number=invoice,
        series="99",
        issue_date=date(2026, 7, 30),
        issuer_name="EMPRESA EMITENTE LTDA",
        issuer_cnpj="46.395.687/0004-55",
        recipient_name="EMPRESA TOMADORA LTDA",
        amount=Decimal(amount),
        quality_flags=() if status is ProcessingStatus.OK else (QualityFlag.MISSING_SERIES,),
        parser_id="nfe-danfe-v1",
    )


def store(path: Path) -> OpenpyxlInvoiceStore:
    return OpenpyxlInvoiceStore(path, now=lambda: NOW)


def test_creates_controlled_workbook_and_hides_technical_columns(tmp_path: Path) -> None:
    path = tmp_path / "Controle_Notas_Fiscais.xlsx"
    result = store(path).upsert(extraction())

    assert result.outcome is UpsertOutcome.CREATED
    workbook = load_workbook(path)
    worksheet = workbook["Notas Fiscais"]
    try:
        assert worksheet.max_row == 2
        assert worksheet["A2"].value == "nf.pdf"
        assert worksheet["C2"].value == "004.241.885"
        assert worksheet["I2"].value == 112000
        assert worksheet["M2"].value == "OK"
        assert worksheet["O2"].value == "a" * 64
        assert worksheet.column_dimensions["O"].hidden is True
        assert worksheet.column_dimensions["R"].hidden is True
    finally:
        workbook.close()


def test_same_sha_updates_same_row_and_preserves_manual_fields(tmp_path: Path) -> None:
    path = tmp_path / "Controle_Notas_Fiscais.xlsx"
    invoice_store = store(path)
    invoice_store.upsert(extraction())

    workbook = load_workbook(path)
    worksheet = workbook["Notas Fiscais"]
    worksheet["J2"] = "OS-5231"
    worksheet["K2"] = date(2026, 12, 31)
    worksheet["L2"] = "Conferido manualmente"
    workbook.save(path)
    workbook.close()

    result = invoice_store.upsert(
        extraction(
            filename="renomeada.pdf",
            invoice="004.241.886",
            amount="113000.50",
        )
    )

    assert result.outcome is UpsertOutcome.UPDATED
    assert result.row == 2
    workbook = load_workbook(path)
    worksheet = workbook["Notas Fiscais"]
    try:
        assert worksheet.max_row == 2
        assert worksheet["A2"].value == "renomeada.pdf"
        assert worksheet["C2"].value == "004.241.886"
        assert worksheet["I2"].value == 113000.5
        assert worksheet["J2"].value == "OS-5231"
        assert worksheet["K2"].value.date() == date(2026, 12, 31)
        assert worksheet["L2"].value == "Conferido manualmente"
    finally:
        workbook.close()


def test_different_sha_creates_new_row(tmp_path: Path) -> None:
    path = tmp_path / "Controle_Notas_Fiscais.xlsx"
    invoice_store = store(path)
    invoice_store.upsert(extraction(sha="a" * 64))

    result = invoice_store.upsert(extraction(sha="b" * 64, filename="nf2.pdf"))

    assert result.outcome is UpsertOutcome.CREATED
    workbook = load_workbook(path)
    worksheet = workbook["Notas Fiscais"]
    try:
        assert worksheet.max_row == 3
    finally:
        workbook.close()


def test_refuses_unknown_existing_workbook_contract(tmp_path: Path) -> None:
    path = tmp_path / "Controle_Notas_Fiscais.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Outra coisa"
    workbook.save(path)
    workbook.close()

    with pytest.raises(WorkbookContractError):
        store(path).upsert(extraction())


def test_atomic_replace_failure_preserves_existing_file_and_cleans_temp(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "Controle_Notas_Fiscais.xlsx"
    invoice_store = store(path)
    invoice_store.upsert(extraction())
    before = path.read_bytes()

    import fiscal_processor.adapters.excel.openpyxl_store as module

    def locked(*_: object) -> None:
        raise PermissionError("locked")

    monkeypatch.setattr(module.os, "replace", locked)

    with pytest.raises(WorkbookLockedError):
        invoice_store.upsert(extraction(invoice="999"))

    assert path.read_bytes() == before
    assert list(tmp_path.glob(".*.tmp.xlsx")) == []
