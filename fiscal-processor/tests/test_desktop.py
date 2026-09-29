from queue import Queue
from threading import Event, Thread

import pytest
from openpyxl import load_workbook

from benchmarks.spatial_corpus import cases, write_pdf
from fiscal_processor.application.batch import BatchResult, DocumentResult, process_batch
from fiscal_processor.presentation.state import row_values, summarize
from fiscal_processor.presentation.worker import run_batch


@pytest.mark.parametrize(
    "batch,state",
    [
        (BatchResult(0, ()), "empty"),
        (BatchResult(2, (DocumentResult(1, "OK"), DocumentResult(2, "FAILED"))), "partial"),
        (BatchResult(1, (DocumentResult(1, "FAILED"),)), "error"),
        (BatchResult(1, (DocumentResult(1, "SKIPPED"),)), "complete"),
        (BatchResult(1, (DocumentResult(1, "OK"),)), "success"),
        (BatchResult(2, (), cancelled=True), "cancelled"),
        (BatchResult(2, (), error_code="WORKBOOK_WRITE_FAILED"), "error"),
    ],
)
def test_summary_preserves_distinct_states(batch, state):
    assert summarize(batch).state == state


def test_skipped_rows_never_invent_previous_values():
    row = row_values("synthetic.pdf", DocumentResult(1, "SKIPPED"))
    assert row == ("synthetic.pdf", "—", "—", "—", "Já registrada")
    assert (
        "não foram reavaliados"
        in summarize(BatchResult(1, (DocumentResult(1, "SKIPPED"),))).message
    )


def test_worker_delivers_saved_real_results_without_ui_access(tmp_path):
    write_pdf(tmp_path / "synthetic.pdf", cases()[0])
    output = tmp_path / "out.xlsx"
    events = Queue()
    worker = Thread(target=run_batch, args=(tmp_path, output, None, False, Event(), events))
    worker.start()
    worker.join(timeout=10)
    assert not worker.is_alive()
    received = []
    while not events.empty():
        received.append(events.get_nowait())
    assert [e.kind for e in received] == ["ready", "started", "document", "done"]
    result = received[2].document
    assert result.outcome == "OK"
    assert result.extraction.invoice_number == cases()[0].expected["number"]
    assert row_values("synthetic.pdf", result)[3].startswith("R$ ")
    book = load_workbook(output)
    assert book.active.cell(2, 3).value == result.extraction.invoice_number
    book.close()


def test_cancel_at_document_boundary_never_starts_next_pdf(tmp_path):
    from fiscal_processor.adapters.excel.openpyxl_store import OpenpyxlInvoiceStore
    from fiscal_processor.composition import build_extractor

    paths = (tmp_path / "a.pdf", tmp_path / "b.pdf")
    for path in paths:
        write_pdf(path, cases()[0])
    cancel = Event()
    starts = []
    result = process_batch(
        paths,
        OpenpyxlInvoiceStore(tmp_path / "out.xlsx"),
        build_extractor(None),
        progress=lambda item, total: cancel.set(),
        cancelled=cancel.is_set,
        started=lambda index, path: starts.append(index),
    )
    assert result.cancelled and len(result.documents) == 1
    assert starts == [1]


def test_worker_empty_and_invalid_folders_are_distinct(tmp_path):
    events = Queue()
    output = tmp_path / "out.xlsx"
    run_batch(tmp_path, output, None, False, Event(), events)
    assert events.get().kind == "ready"
    assert summarize(events.get().batch).state == "empty"
    assert not output.exists()
    run_batch(tmp_path / "absent", output, None, False, Event(), events)
    assert events.get().kind == "error"
