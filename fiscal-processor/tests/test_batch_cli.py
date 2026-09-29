from openpyxl import load_workbook

from benchmarks.spatial_corpus import cases, write_pdf
from fiscal_processor.adapters.excel.openpyxl_store import OpenpyxlInvoiceStore
from fiscal_processor.application.batch import discover_pdfs, process_batch
from fiscal_processor.cli import main
from fiscal_processor.domain import DocumentType, ExtractionMode, FiscalExtraction, ProcessingStatus


def test_real_cli_deduplicates_and_preserves_manual_fields(tmp_path, capsys):
    source = tmp_path / "input"
    source.mkdir()
    write_pdf(source / "a.pdf", cases()[0])
    (source / "b.PDF").write_bytes((source / "a.pdf").read_bytes())
    output = tmp_path / "output.xlsx"
    args = [str(source), "--output", str(output)]
    assert main(args) == 0
    assert "SKIPPED=1" in capsys.readouterr().out
    book = load_workbook(output)
    sheet = book.active
    assert sheet.max_row == 2
    assert sheet.cell(2, 3).value == cases()[0].expected["number"]
    sheet.cell(2, 10).value = "OS manual"
    sheet.cell(2, 11).value = "validade manual"
    sheet.cell(2, 12).value = "observação manual"
    book.save(output)
    book.close()
    original = output.read_bytes()
    assert main(args) == 0
    assert output.read_bytes() == original  # skip does not save or change timestamps
    assert main(args + ["--reprocess"]) == 0
    book = load_workbook(output)
    assert book.active.max_row == 2
    assert [book.active.cell(2, c).value for c in (10, 11, 12)] == [
        "OS manual",
        "validade manual",
        "observação manual",
    ]
    book.close()


def test_corrupt_pdf_does_not_stop_next_document_or_leak_name(tmp_path, capsys):
    (tmp_path / "a-secret.pdf").write_bytes(b"bad secret fiscal content")
    write_pdf(tmp_path / "b.pdf", cases()[1])
    output = tmp_path / "out.xlsx"
    assert main([str(tmp_path), "--output", str(output)]) == 1
    log = capsys.readouterr().out
    assert "FAILED=1" in log and "OK=1" in log
    assert "secret" not in log and str(tmp_path) not in log
    book = load_workbook(output)
    assert book.active.max_row == 2
    book.close()


def test_empty_folder_does_not_create_workbook(tmp_path, capsys):
    output = tmp_path / "out.xlsx"
    assert main([str(tmp_path), "--output", str(output)]) == 0
    assert "EMPTY" in capsys.readouterr().out
    assert not output.exists()


def test_discovery_is_nonrecursive_case_insensitive_and_ignores_links(tmp_path):
    (tmp_path / "B.PDF").write_bytes(b"b")
    (tmp_path / "a.pdf").write_bytes(b"a")
    (tmp_path / "x.txt").write_bytes(b"x")
    child = tmp_path / "nested"
    child.mkdir()
    (child / "c.pdf").write_bytes(b"c")
    assert [p.name for p in discover_pdfs(tmp_path)] == ["a.pdf", "B.PDF"]


def extraction(path, digest):
    return FiscalExtraction(
        digest, path.name, DocumentType.UNKNOWN, ExtractionMode.NATIVE_TEXT, ProcessingStatus.REVIEW
    )


def test_failed_reprocessing_leaves_existing_workbook_untouched(tmp_path):
    path = tmp_path / "a.pdf"
    path.write_bytes(b"synthetic")
    output = tmp_path / "out.xlsx"
    store = OpenpyxlInvoiceStore(output)
    assert process_batch((path,), store, extraction).documents[0].outcome == "REVIEW"
    original = output.read_bytes()

    def fail(path, digest):
        raise RuntimeError("private fiscal text")

    result = process_batch((path,), store, fail, reprocess=True)
    assert result.documents[0].error_code == "EXTRACTION_FAILED"
    assert output.read_bytes() == original


def test_changed_input_never_persists_wrong_hash(tmp_path):
    path = tmp_path / "a.pdf"
    path.write_bytes(b"before")
    output = tmp_path / "out.xlsx"

    def mutate(path, digest):
        path.write_bytes(b"after")
        return extraction(path, digest)

    result = process_batch((path,), OpenpyxlInvoiceStore(output), mutate)
    assert result.documents[0].outcome == "FAILED"
    assert not output.exists()


def test_workbook_failure_stops_batch_without_claiming_saved_results(tmp_path, monkeypatch):
    paths = (tmp_path / "a.pdf", tmp_path / "b.pdf")
    for path in paths:
        path.write_bytes(path.name.encode())
    store = OpenpyxlInvoiceStore(tmp_path / "out.xlsx")
    seen = []

    def locked(book):
        raise PermissionError("private path")

    monkeypatch.setattr(store, "_atomic_save", locked)
    result = process_batch(paths, store, extraction, progress=lambda item, total: seen.append(item))
    assert result.error_code == "WORKBOOK_WRITE_FAILED"
    assert result.total == 2 and len(result.documents) == 1
    assert seen[0].outcome == "FAILED"
    assert not store.path.exists()


def test_invalid_workbook_is_preserved_and_cli_reports_fatal(tmp_path, capsys):
    write_pdf(tmp_path / "a.pdf", cases()[0])
    output = tmp_path / "out.xlsx"
    output.write_bytes(b"not a workbook")
    assert main([str(tmp_path), "--output", str(output)]) == 2
    assert output.read_bytes() == b"not a workbook"
    assert "WORKBOOK_READ_FAILED" in capsys.readouterr().out


def test_missing_ocr_configuration_fails_only_blank_document(tmp_path, capsys):
    import pypdfium2

    with pypdfium2.PdfDocument.new() as document:
        page = document.new_page(612, 792)
        page.close()
        document.save(tmp_path / "a-blank.pdf")
    write_pdf(tmp_path / "b-native.pdf", cases()[0])
    output = tmp_path / "out.xlsx"
    assert main([str(tmp_path), "--output", str(output)]) == 1
    log = capsys.readouterr().out
    assert "FAILED=1" in log and "OK=1" in log


def test_cancel_propagates_and_preserves_completed_documents(tmp_path):
    import pytest

    paths = (tmp_path / "a.pdf", tmp_path / "b.pdf")
    for path in paths:
        path.write_bytes(path.name.encode())
    output = tmp_path / "out.xlsx"

    def interrupt(path, digest):
        if path == paths[1]:
            raise KeyboardInterrupt
        return extraction(path, digest)

    with pytest.raises(KeyboardInterrupt):
        process_batch(paths, OpenpyxlInvoiceStore(output), interrupt)
    book = load_workbook(output)
    assert book.active.max_row == 2
    assert book.active.cell(2, 1).value == "a.pdf"
    book.close()
