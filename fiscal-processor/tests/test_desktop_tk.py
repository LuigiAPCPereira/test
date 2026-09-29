"""Real Tk controls and event loop. CI must set FISCAL_REQUIRE_DISPLAY=1."""

import os
import time
from threading import Event

import pytest
from openpyxl import load_workbook

from benchmarks.spatial_corpus import cases, write_pdf
from fiscal_processor.application.batch import BatchResult, DocumentResult
from fiscal_processor.presentation.worker import WorkerEvent


@pytest.fixture
def desktop():
    try:
        import tkinter as tk

        root = tk.Tk()
    except Exception as exc:
        # TclError isn't available when the import itself fails.
        if os.environ.get("FISCAL_REQUIRE_DISPLAY") == "1":
            pytest.fail(f"required Tk display unavailable: {type(exc).__name__}")
        pytest.skip("Tk display unavailable")
    from fiscal_processor.presentation.desktop import Desktop

    app = Desktop(root)
    root.update()
    root.focus_force()
    root.update()
    try:
        yield app
    finally:
        app.cancel.set()
        root.destroy()


def select(app, folder, monkeypatch):
    monkeypatch.setattr("tkinter.filedialog.askdirectory", lambda **kwargs: str(folder))
    app.choose.invoke()
    app.root.update()


def drain_until_complete(app):
    deadline = time.monotonic() + 10
    while app.busy and time.monotonic() < deadline:
        app.root.update()
        Event().wait(0.01)
    assert not app.busy, "worker did not finish within 10 seconds"
    app.root.update()


def test_initial_selection_and_keyboard_controls(desktop, tmp_path, monkeypatch):
    app = desktop
    assert str(app.process["state"]) == "disabled"
    assert str(app.open_button["state"]) == "disabled"
    select(app, tmp_path, monkeypatch)
    assert str(app.process["state"]) == "normal"
    app.root.focus_force()
    app.choose.focus_set()
    app.root.update()
    app.choose.event_generate("<Tab>")
    app.root.update()
    assert app.root.focus_get() == app.process


def test_buttons_process_real_pdf_and_table_matches_workbook(desktop, tmp_path, monkeypatch):
    write_pdf(tmp_path / "nota-sintetica.pdf", cases()[0])
    app = desktop
    select(app, tmp_path, monkeypatch)
    app.process.invoke()
    assert app.busy
    assert str(app.choose["state"]) == "disabled"
    assert str(app.open_button["state"]) == "disabled"
    drain_until_complete(app)
    children = app.table.get_children()
    assert len(children) == 1
    values = app.table.item(children[0], "values")
    assert values[2] == cases()[0].expected["number"]
    assert values[4] == "Processada"
    assert str(app.open_button["state"]) == "normal"
    book = load_workbook(tmp_path / "Controle_Notas_Fiscais.xlsx")
    assert book.active.cell(2, 3).value == values[2]
    book.close()
    assert app.root.focus_get() == app.process


def test_empty_folder_is_not_error_or_fake_completion(desktop, tmp_path, monkeypatch):
    select(desktop, tmp_path, monkeypatch)
    desktop.process.invoke()
    drain_until_complete(desktop)
    assert "Nenhum PDF" in desktop.status.get()
    assert not desktop.table.get_children()
    assert str(desktop.open_button["state"]) == "disabled"


def test_cancel_and_close_wait_for_worker(desktop):
    app = desktop
    app.busy = True
    app.close()
    assert app.cancel.is_set()
    assert app.root.winfo_exists()
    assert "Aguardando" in app.status.get()
    app.handle(WorkerEvent("done", batch=BatchResult(2, (), cancelled=True)))
    assert not app.busy
    assert "Cancelado" in app.status.get()


def test_partial_failure_keeps_rows_and_workbook_guidance(desktop):
    app = desktop
    app.handle(WorkerEvent("ready", total=2))
    app.handle(
        WorkerEvent(
            "document",
            total=2,
            filename="nome-longo-" * 20 + ".pdf",
            document=DocumentResult(1, "SKIPPED"),
        )
    )
    failed = DocumentResult(2, "FAILED", error_code="WORKBOOK_WRITE_FAILED")
    app.handle(WorkerEvent("document", total=2, filename="synthetic.pdf", document=failed))
    app.handle(WorkerEvent("done", batch=BatchResult(2, (failed,), "WORKBOOK_WRITE_FAILED")))
    assert len(app.table.get_children()) == 2
    key = app.table.get_children()[-1]
    app.table.selection_set(key)
    app.root.update()
    assert "Feche o Excel" in app.current.get()
    assert "preservados" in app.status.get()
