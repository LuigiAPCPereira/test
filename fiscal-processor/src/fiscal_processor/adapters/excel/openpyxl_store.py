"""Idempotent XLSX store with explicit ownership of manual columns."""

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from datetime import UTC, datetime
from enum import StrEnum
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fiscal_processor.domain.invoice import FiscalExtraction


class ExcelStoreError(RuntimeError):
    """Base error for workbook persistence failures."""


class WorkbookContractError(ExcelStoreError):
    """Existing workbook does not match the controlled column contract."""


class WorkbookLockedError(ExcelStoreError):
    """Workbook could not be replaced, usually because it is open/locked."""


class UpsertOutcome(StrEnum):
    CREATED = "CREATED"
    UPDATED = "UPDATED"


@dataclass(frozen=True, slots=True)
class UpsertResult:
    outcome: UpsertOutcome
    row: int


class OpenpyxlInvoiceStore:
    SHEET_NAME = "Notas Fiscais"
    HEADERS = (
        "Arquivo",
        "Tipo",
        "Número da NF",
        "Série",
        "Data de emissão",
        "Empresa prestadora/emitente",
        "CNPJ",
        "Empresa tomadora/destinatária",
        "Valor",
        "Número da OS",
        "Validade",
        "Observações",
        "Situação",
        "Motivo da revisão",
        "_sha256",
        "_extraction_mode",
        "_parser_id",
        "_processed_at",
    )
    MANUAL_COLUMNS = {10, 11, 12}
    SHA_COLUMN = 15

    def __init__(
        self,
        path: str | Path,
        *,
        now: Callable[[], datetime] | None = None,
    ) -> None:
        self.path = Path(path)
        self._now = now or (lambda: datetime.now(UTC))

    def upsert(self, extraction: FiscalExtraction) -> UpsertResult:
        workbook, worksheet = self._load_or_create()
        try:
            matching_rows = [
                row
                for row in range(2, worksheet.max_row + 1)
                if worksheet.cell(row, self.SHA_COLUMN).value == extraction.source_sha256
            ]
            if len(matching_rows) > 1:
                raise WorkbookContractError(
                    f"duplicate technical key {extraction.source_sha256} found in workbook"
                )

            if matching_rows:
                row = matching_rows[0]
                outcome = UpsertOutcome.UPDATED
            else:
                row = max(worksheet.max_row + 1, 2)
                outcome = UpsertOutcome.CREATED

            self._write_automatic_fields(worksheet, row, extraction)
            self._atomic_save(workbook)
            return UpsertResult(outcome=outcome, row=row)
        finally:
            workbook.close()

    def _load_or_create(self) -> tuple[Any, Worksheet]:
        if self.path.exists():
            try:
                workbook = load_workbook(self.path)
            except Exception as exc:
                raise ExcelStoreError(f"unable to open workbook: {self.path.name}") from exc
            if self.SHEET_NAME not in workbook.sheetnames:
                workbook.close()
                raise WorkbookContractError(
                    f"worksheet {self.SHEET_NAME!r} not found in existing workbook"
                )
            worksheet = workbook[self.SHEET_NAME]
            self._validate_headers(worksheet)
            return workbook, worksheet

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.SHEET_NAME
        self._initialize_sheet(worksheet)
        return workbook, worksheet

    def _initialize_sheet(self, worksheet: Worksheet) -> None:
        worksheet.append(self.HEADERS)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = "A1:N1"
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for column in range(15, len(self.HEADERS) + 1):
            worksheet.column_dimensions[get_column_letter(column)].hidden = True
        worksheet.column_dimensions["A"].width = 28
        worksheet.column_dimensions["F"].width = 34
        worksheet.column_dimensions["H"].width = 34
        worksheet.column_dimensions["L"].width = 36
        worksheet.column_dimensions["N"].width = 32

    def _validate_headers(self, worksheet: Worksheet) -> None:
        observed = tuple(
            worksheet.cell(1, column).value
            for column in range(1, len(self.HEADERS) + 1)
        )
        if observed != self.HEADERS:
            raise WorkbookContractError("workbook headers do not match the controlled contract")

    def _write_automatic_fields(
        self,
        worksheet: Worksheet,
        row: int,
        extraction: FiscalExtraction,
    ) -> None:
        reason = ", ".join(flag.value for flag in extraction.quality_flags) or None
        processed_at = self._now()
        if processed_at.tzinfo is None:
            processed_at = processed_at.replace(tzinfo=UTC)
        else:
            processed_at = processed_at.astimezone(UTC)

        values = {
            1: extraction.source_filename,
            2: extraction.document_type.value,
            3: extraction.invoice_number,
            4: extraction.series,
            5: extraction.issue_date,
            6: extraction.issuer_name,
            7: extraction.issuer_cnpj,
            8: extraction.recipient_name,
            9: float(extraction.amount) if extraction.amount is not None else None,
            13: extraction.status.value,
            14: reason,
            15: extraction.source_sha256,
            16: extraction.extraction_mode.value,
            17: extraction.parser_id,
            18: processed_at.isoformat(),
        }
        for column, value in values.items():
            if column in self.MANUAL_COLUMNS:
                raise AssertionError("automatic writer must never own manual columns")
            worksheet.cell(row, column).value = value

        worksheet.cell(row, 3).number_format = "@"
        worksheet.cell(row, 4).number_format = "@"
        worksheet.cell(row, 5).number_format = "DD/MM/YYYY"
        worksheet.cell(row, 7).number_format = "@"
        worksheet.cell(row, 9).number_format = "#,##0.00"
        worksheet.cell(row, 15).number_format = "@"

    def _atomic_save(self, workbook: Any) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        fd, temporary_name = tempfile.mkstemp(
            prefix=f".{self.path.stem}.",
            suffix=".tmp.xlsx",
            dir=self.path.parent,
        )
        os.close(fd)
        temporary = Path(temporary_name)
        try:
            workbook.save(temporary)
            try:
                os.replace(temporary, self.path)
            except PermissionError as exc:
                raise WorkbookLockedError(
                    f"unable to replace workbook {self.path.name}; it may be open or locked"
                ) from exc
            except OSError as exc:
                raise ExcelStoreError(f"unable to replace workbook {self.path.name}") from exc
        finally:
            temporary.unlink(missing_ok=True)
